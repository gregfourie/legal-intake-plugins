#!/usr/bin/env python3
"""
build_index.py — matter-index engine (v2.5.1, 26 September 2026).

Walks a matter folder, detects metadata by HEURISTIC (no AI), and writes:

    OUTPUT_PATH/index.md          one structured entry per document
    OUTPUT_PATH/docs/NNNN_*.md    extracted text ONLY for files with no intake mirror
    OUTPUT_PATH/_illegible.md     scanned / failed / stale-mirror / password files
    OUTPUT_PATH/_state.json       SHA-256 incremental cache

v2.5 (25 Sep 2026, burn-night D1 — proposed):
  * Text-poor mirrors, second signal. chars/page < 200 stays the first test.
    A mirror also counts as text-poor when it has fewer than
    --text-poor-lines (default 2) non-empty text lines per physical page, or
    when (over 3+ pages) 70% or more of its text by characters sits in
    lines repeated three times or more — the gazette pattern: a running
    header on every page and no body.
  * Probable versions. After indexing, entries are compared by normalised
    title (filename stem with version words, dates and numbers removed) and
    by opening text. Pairs that look like the same document in another
    file or format (docx/pdf pair, two scans of one judgment, draft v
    signed) are listed under "## Probable versions" in index.md — a
    reading list for the map readers, never a statement of fact.
  * Phase 2 gate. On a mirror-first matter (or with --gate) the run ends
    with "GATE: PASS" or "GATE: FAIL" and exit code 2 on failure. The gate
    fails when no filed document was indexed or any text-poor mirror was
    found. The index is still written so the gaps are on record; the
    orchestrator stops and reports. --no-gate switches the gate off.
  * --include-hopper-input refuses to run (exit 3) when the matter path
    lies under a "live cases" folder. It is a workflow-test flag only.

v2.5.1 (26 Sep 2026 — proposed): DOCX-type mirrors. Paragraph markers
  `<!-- PARA n -->` and `<!-- PARA n | 3.2 -->` (own numbering) are counted
  and the entry says "cite as (docx, para n)"; PARA, TABLE and FOOTNOTES
  markers are stripped before the first substantive paragraph is picked
  (a marker on its own line directly above the text no longer hides the
  paragraph). Ported from the parallel D1 package (engine v2.4, 25 Sep
  2026) and widened to the pipe form of the proposed AGENTS.md standard.

v2.3 (15 Sep 2026): page markers may be `<!-- PHYSICAL PAGE n -->` or the
  `## Physical page n` heading form; when only `.ai/document-state.json`
  exists, run `_tools/make_mirror_index.py` first to build mirror-index.json.

v2 changes
  * Mirror-first. If the matter has an intake mirror index
    (default .ai/mirror-index.json), a file whose SHA-256 matches a row's
    source_sha256 or derivative_sha256 takes its text from the recorded
    mirror (.ai/mirrors/<hash>.md). No re-extraction, no docs/ copy; the
    index entry points at the mirror. A mirror whose on-disk hash differs
    from the recorded mirror_sha256 is STALE: flagged, never re-extracted
    silently.
  * Exclusions. Skips any path component starting with "_" or "."
    (so .ai, _index, _build ...), the hopper, archive and audit folders,
    volume splits and superseded drafts, and the matter control files
    (MATTER.md, _matter.md, LESSONS.md, AGENTS.md).
  * Duplicates. A file whose hash was already indexed (e.g. the active copy
    in "00 Current") is recorded as a duplicate of the first entry, not
    indexed twice.

Usage:
    python3 build_index.py MATTER_PATH OUTPUT_PATH [--no-recursive] [--force]
        [--mirror-index PATH] [--exclude-dir NAME ...] [--exclude-file NAME ...]
        [--text-poor-lines N] [--gate | --no-gate] [--no-versions]

Exit codes: 0 pass; 2 phase 2 gate failed (index written); 3 refused
(--include-hopper-input on a live matter).

This script never modifies source files or mirrors.
"""

import argparse
import difflib
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".xlsx", ".msg", ".eml", ".txt", ".md", ".rtf"}

EXCLUDE_DIRS = {
    "01 Hopper", "05 Archive", "06 Audits",
    "Bookmarked PDFs", "legibility-review", "Previous Drafts", "eml",
    "needs adobe ocr",
}
EXCLUDE_FILES = {"MATTER.md", "_matter.md", "LESSONS.md", "AGENTS.md",
                 "START HERE.md", "Index.md", "Weekly Updates.md"}

MONTHS = ("january|february|march|april|may|june|july|august|september|"
          "october|november|december")
DATE_PATTERNS = [
    (re.compile(r"\b(20\d{2}|19\d{2})-(\d{2})-(\d{2})\b"), "ymd"),
    (re.compile(rf"\b(\d{{1,2}})\s+({MONTHS})\s+(20\d{{2}}|19\d{{2}})\b",
                re.IGNORECASE), "dmy"),
    (re.compile(rf"\b({MONTHS})\s+(\d{{1,2}}),?\s+(20\d{{2}})\b",
                re.IGNORECASE), "mdy"),
]

TYPE_SIGNALS = [
    ("affidavit", "high", re.compile(
        r"do hereby make oath and (say|state)|states? under oath|"
        r"I,? the undersigned", re.IGNORECASE)),
    ("judgment", "high", re.compile(
        r"J\s?U\s?D\s?G\s?M\s?E\s?N\s?T|JUDGMENT\b.{0,40}\n", )),
    ("court order", "high", re.compile(
        r"IT IS ORDERED|THE FOLLOWING ORDER IS MADE|ORDER OF COURT",
        re.IGNORECASE)),
    ("notice of motion", "high", re.compile(r"NOTICE OF MOTION",
                                            re.IGNORECASE)),
    ("heads of argument", "high", re.compile(
        r"HEADS OF ARGUMENT|WRITTEN SUBMISSIONS", re.IGNORECASE)),
    ("pleading", "medium", re.compile(
        r"PARTICULARS OF CLAIM|STATEMENT OF (CLAIM|DEFENCE)|REPLICATION|"
        r"\bPLEA\b|NOTICE OF (APPEAL|APPLICATION)", re.IGNORECASE)),
    ("award", "medium", re.compile(r"ARBITRATION AWARD|\bAWARD\b.{0,30}\n",
                                   re.IGNORECASE)),
    ("letter", "medium", re.compile(
        r"\bDear\s+(Sir|Madam|Mr|Ms|Mrs|Adv|Judge|Colleagues?)\b|"
        r"Yours faithfully|Yours sincerely", re.IGNORECASE)),
    ("email", "medium", re.compile(r"^From:.*\n(To|Sent|Date):", re.MULTILINE)),
]

# Generic defaults. A jurisdiction profile (_intake-profile.json, written by
# the intake-setup skill) adds its own court names and case-number patterns.
TRIBUNALS = [
    "SUPREME COURT", "COURT OF APPEAL", "CONSTITUTIONAL COURT", "HIGH COURT",
    "DISTRICT COURT", "COUNTY COURT", "MAGISTRATE", "CIRCUIT COURT",
    "EMPLOYMENT TRIBUNAL", "TRIBUNAL", "ARBITRATION",
]

CASE_REFS = [re.compile(
    r"\bCASE\s*(?:NO|NUMBER)[.:\s]*([A-Z]{0,4}\s?[\w-]*\d[\w/.-]*)",
    re.IGNORECASE)]


def find_profile(start, explicit=None):
    """Return the parsed jurisdiction profile, or {} if none.
    Looks for --profile, else _intake-profile.json in start or its parents."""
    if explicit:
        return json.loads(Path(explicit).read_text())
    for d in [Path(start).resolve(), *Path(start).resolve().parents]:
        f = d / "_intake-profile.json"
        if f.is_file():
            return json.loads(f.read_text())
    return {}


def apply_profile(prof):
    global TRIBUNALS, CASE_REFS
    t = [x.upper() for x in prof.get("tribunals", [])]
    # longest first, so "LABOUR APPEAL COURT" wins over "LABOUR COURT"
    TRIBUNALS = sorted(set(t) | set(TRIBUNALS), key=len, reverse=True)
    CASE_REFS = [re.compile(r, re.IGNORECASE)
                 for r in prof.get("case_ref_patterns", [])] + CASE_REFS

PAGE_MARKER = re.compile(r"(?:<!--\s*PHYSICAL(?:\s+PDF)?\s+PAGE\s+(\d+)|^## Physical page (\d+)\s*$)", re.IGNORECASE | re.MULTILINE)  # v2.3: heading form written by phase-one-intake (Codex) accepted too
END_PROVENANCE = re.compile(r"<!--\s*END CURRENT PROVENANCE\s*-->", re.IGNORECASE)
# v2.5.1: DOCX-type mirrors carry paragraph markers, not page markers.
PARA_MARKER = re.compile(r"<!--\s*PARA\s+(\d+)(?:\s*\|[^>]*?)?\s*-->", re.IGNORECASE)
BODY_MARKER = re.compile(
    r"<!--\s*(?:PARA\s+\d+(?:\s*\|[^>]*?)?|TABLE\s+\d+|FOOTNOTES)\s*-->", re.IGNORECASE)


def sha256(path, blocksize=1 << 20):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(blocksize)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def safe_name(stem):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_").lower()[:60]


# ---------------------------------------------------------------- extract

def extract_pdf(path):
    import pdfplumber
    pages, total_chars = [], 0
    with pdfplumber.open(path) as pdf:
        n_pages = len(pdf.pages)
        for pg in pdf.pages:
            t = pg.extract_text() or ""
            pages.append(t)
            total_chars += len(t)
    text = "\n\n".join(pages)
    avg = total_chars / max(n_pages, 1)
    status = "good" if avg >= 200 else ("partial" if avg >= 50 else "scanned")
    return text, n_pages, status


def extract_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    text = "\n".join(parts)
    return text, None, "good" if text.strip() else "empty"


def extract_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title} ({ws.max_row} rows x "
                     f"{ws.max_column} cols)")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 5:
                parts.append("... (structure only; open source for full data)")
                break
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(parts), None, "good"


def extract_msg(path):
    import extract_msg
    m = extract_msg.Message(str(path))
    header = (f"From: {m.sender}\nTo: {m.to}\nDate: {m.date}\n"
              f"Subject: {m.subject}\n"
              f"Attachments: {', '.join(a.longFilename or a.shortFilename or '?' for a in m.attachments) or 'none'}\n")
    return header + "\n" + (m.body or ""), None, "good"


def extract_eml(path):
    import email
    from email import policy
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    body_text = body.get_content() if body else ""
    atts = [p.get_filename() for p in msg.iter_attachments()]
    header = (f"From: {msg.get('From')}\nTo: {msg.get('To')}\n"
              f"Date: {msg.get('Date')}\nSubject: {msg.get('Subject')}\n"
              f"Attachments: {', '.join(a for a in atts if a) or 'none'}\n")
    return header + "\n" + body_text, None, "good"


def extract_text_file(path):
    text = Path(path).read_text(errors="replace")
    return text, None, "good"


def extract_rtf(path):
    # Plain-text fallback: strip RTF control words. Prefer an intake mirror.
    raw = Path(path).read_text(errors="replace")
    text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)
    return text, None, "partial"


EXTRACTORS = {
    ".pdf": extract_pdf, ".docx": extract_docx, ".xlsx": extract_xlsx,
    ".msg": extract_msg, ".eml": extract_eml,
    ".txt": extract_text_file, ".md": extract_text_file, ".rtf": extract_rtf,
}


# ----------------------------------------------------------------- mirrors

def load_mirror_index(matter, path_arg):
    p = Path(path_arg) if path_arg else matter / ".ai" / "mirror-index.json"
    if not p.is_absolute():
        p = matter / p
    if not p.exists():
        return {}, None
    data = json.loads(p.read_text())
    rows = data.get("rows", data if isinstance(data, list) else [])
    by_hash = {}
    for r in rows:
        for key in ("source_sha256", "derivative_sha256"):
            h = r.get(key)
            if h and r.get("mirror_path"):
                by_hash.setdefault(h, r)
    return by_hash, p


TEXT_POOR_LINES_PER_PAGE = 2      # second signal (a): fewer text lines per page than this
TEXT_POOR_REPEAT_SHARE = 0.70     # second signal (b): share of text in lines repeated 3+ times


def text_poor_second_signal(text_only, n_pages, min_lines=TEXT_POOR_LINES_PER_PAGE):
    """Gazette pattern: a running header on every page and no body.

    Two tests, either one flags: (a) fewer than min_lines non-empty text
    lines per physical page; (b) over 3+ pages, at least 70% of the text
    (by characters) sits in lines that recur three times or more — a header
    or footer block, not body text. Body text almost never repeats verbatim,
    so a judgment with long paragraph lines passes (b) even when it has few
    lines per page. Returns a reason string or None. The caller has already
    stripped page markers and comments.
    """
    if not n_pages:
        return None
    lines = [" ".join(l.split()) for l in text_only.splitlines()]
    lines = [l for l in lines if l]
    if not lines:
        return f"0 text lines over {n_pages} pages"
    per_page = len(lines) / n_pages
    if per_page < min_lines:
        return (f"{len(lines)} text lines over {n_pages} pages "
                f"({per_page:.1f}/page, < {min_lines}) — headers only?")
    if n_pages >= 3:
        counts = {}
        for l in lines:
            k = l.lower()
            counts[k] = counts.get(k, 0) + 1
        total = sum(len(l) for l in lines)
        repeated = sum(len(l) for l in lines if counts[l.lower()] >= 3)
        share = repeated / total if total else 0
        if share >= TEXT_POOR_REPEAT_SHARE:
            return (f"{share:.0%} of the text sits in lines repeated 3+ times "
                    f"over {n_pages} pages — running headers, no body?")
    return None


def read_mirror(matter, row, min_lines=TEXT_POOR_LINES_PER_PAGE):
    """Return (text_without_header, n_pages, status, problem)."""
    mp = matter / row["mirror_path"]
    if not mp.exists():
        return None, None, "failed", f"mirror missing: {row['mirror_path']}"
    on_disk = sha256(mp)
    recorded = row.get("mirror_sha256")
    if recorded and on_disk != recorded:
        return None, None, "failed", (
            f"STALE MIRROR — on-disk hash {on_disk[:12]} != recorded "
            f"{recorded[:12]}; refresh through intake, do not re-extract here")
    raw = mp.read_text(errors="replace")
    # Strip the mirror's provenance header: the document text starts at the
    # first physical-page marker. Fallbacks: after an END CURRENT PROVENANCE
    # comment, after a YAML front-matter block, else the whole file.
    body = raw
    m = PAGE_MARKER.search(raw)
    if m:
        body = raw[m.start():]
    else:
        e = END_PROVENANCE.search(raw)
        if e:
            body = raw[e.end():]
        elif raw.startswith("---"):
            end = raw.find("\n---", 3)
            if end != -1:
                body = raw[end + 4:]
    markers = [a or b for a, b in PAGE_MARKER.findall(body)]
    n_pages = max((int(m) for m in markers), default=None) or row.get("page_count")
    status = "mirror-verified" if row.get("verified") else (
        "mirror-read" if row.get("complete_reading") else "mirror-extracted")
    # A mirror can exist and still hold almost no text (un-OCR'd scan with a
    # preliminary extraction). Flag it; do not treat it as readable.
    text_only = PAGE_MARKER.sub("", body)
    text_only = re.sub(r"<!--.*?-->", "", text_only, flags=re.S)
    chars = len(text_only.strip())
    if n_pages and chars / n_pages < 200:
        return body, n_pages, "mirror-text-poor", (
            f"TEXT-POOR MIRROR — {chars} chars over {n_pages} pages "
            f"(< 200/page); scan not OCR'd or extraction failed; route through intake OCR")
    # v2.5 second signal: enough characters per page, but only a running
    # header on each page (gazette pattern). Flagged the same way.
    why = text_poor_second_signal(text_only, n_pages, min_lines)
    if why:
        return body, n_pages, "mirror-text-poor", (
            f"TEXT-POOR MIRROR (second signal) — {why}; "
            f"read the source PDF before relying on it; route through intake OCR if scanned")
    return body, n_pages, status, None


# --------------------------------------------------------------- metadata

def detect_type(text, ext):
    head = text[:6000]
    for typ, conf, rx in TYPE_SIGNALS:
        if rx.search(head):
            return typ, conf
    if ext == ".xlsx":
        return "spreadsheet", "high"
    if ext in (".msg", ".eml"):
        return "email", "high"
    return "document", "low"


def detect_date(text, filename):
    head = text[:8000]
    for rx, kind in DATE_PATTERNS:
        m = rx.search(head)
        if m:
            try:
                if kind == "ymd":
                    d = datetime(int(m.group(1)), int(m.group(2)),
                                 int(m.group(3)))
                elif kind == "dmy":
                    d = datetime.strptime(
                        f"{m.group(1)} {m.group(2)[:3].title()} {m.group(3)}",
                        "%d %b %Y")
                else:
                    d = datetime.strptime(
                        f"{m.group(2)} {m.group(1)[:3].title()} {m.group(3)}",
                        "%d %b %Y")
                return d.strftime("%Y-%m-%d"), "high — explicit date in document"
            except ValueError:
                continue
    m = re.search(r"(20\d{2})[-_. ]?(\d{2})[-_. ]?(\d{2})", filename)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}", "low — from filename"
    return "—", "none found"


def detect_parties(text):
    m = re.search(
        r"[Ii]n the matter between[:\s]*\n(.{1,200}?)\n.*?\b[Aa]nd\b\s*\n(.{1,200}?)\n",
        text[:5000], re.DOTALL)
    if m:
        a = " ".join(m.group(1).split())[:80]
        b = " ".join(m.group(2).split())[:80]
        return f"{a} / {b}", "medium"
    return "—", "none found"


def detect_deponent(text):
    m = re.search(r"I,? the undersigned,?\s*\n?\s*([A-Z][A-Za-z .'-]{3,60})",
                  text[:5000])
    if m:
        return m.group(1).strip().rstrip(","), "medium"
    return "—", "none found"


def detect_caseref(text):
    for rx in CASE_REFS:
        m = rx.search(text[:5000])
        if m:
            return (m.group(1) if m.groups() else m.group(0)).strip(), "medium"
    return "—", "none found"


def detect_tribunal(text):
    head = text[:5000].upper()
    for t in TRIBUNALS:
        if t in head:
            return t.title(), "medium"
    return "—", "none found"


def first_substantive_para(text):
    for para in text.split("\n\n"):
        p = " ".join(BODY_MARKER.sub("", para).split())   # v2.5.1: a marker is not the paragraph
        if p.startswith("<!--"):
            continue
        if len(p) > 100:
            return p[:400]
    return ""


def count_numbered_paras(text):
    return len(re.findall(r"^\s*\d{1,3}[.)]\s+\S", text, re.MULTILINE))


# ------------------------------------------------------ probable versions

VERSION_WORDS = re.compile(
    r"\b(v\d+|version|draft|final|signed|amended|revised|clean|marked|redline|"
    r"tracked|copy|scan|scanned|ocr|ocrd|searchable|filed|stamped|as filed|"
    r"as sent|sent|unsigned|latest|new|old|updated|current|original|generated)\b",
    re.IGNORECASE)
DATE_TOKEN = re.compile(r"\b(20|19)\d{2}[-_. ]?\d{2}[-_. ]?\d{2}\b|\(?\b(20|19)\d{2}\)?")
LEADING_SEQ = re.compile(r"^[a-z]?-?\d{1,4}[-_. )]+", re.IGNORECASE)


def norm_title(stem):
    """Filename stem without version words, dates, years and a leading
    sequence number. Statutory and case-number digits are KEPT (so
    "Section 198B(10)" and "Section 198B(6)" stay different)."""
    t = stem.lower()
    t = LEADING_SEQ.sub(" ", t)
    t = DATE_TOKEN.sub(" ", t)
    t = VERSION_WORDS.sub(" ", t)
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return " ".join(t.split())


def digit_tokens(norm):
    return sorted(re.findall(r"\d+", norm))


def norm_opening(first_para):
    t = re.sub(r"<!--.*?-->", "", first_para or "")
    t = re.sub(r"[^a-z ]+", " ", t.lower())
    return " ".join(t.split())[:300]


def probable_versions(entries, title_cut=0.85, opening_cut=0.90, max_entries=3000):
    """Pairs of entries that look like the same document in another file or
    format. Heuristic only — a reading list for the map readers.

    Title test: normalised-title similarity >= title_cut AND the same digit
    tokens (a statutory series or numbered set is not a version family).
    Opening test: first-paragraph similarity >= opening_cut, but an opening
    shared by three or more entries is treated as boilerplate (a publisher
    or commentary header) and reported once, not pairwise.

    Returns (pairs, notes). pairs: list of (seq_a, seq_b, reason)."""
    notes = []
    if len(entries) > max_entries:
        return [], [f"probable-versions pass skipped: {len(entries)} entries "
                    f"exceeds the {max_entries} cap; run doc_profile.py per folder"]
    prepared = []
    for e in entries:
        stem = Path(e["source"]).stem
        nt = norm_title(stem)
        prepared.append((e, nt, digit_tokens(nt), norm_opening(e.get("first_para", "")),
                         Path(e["source"]).suffix.lower()))
    # boilerplate openings: identical normalised opening on 3+ entries
    groups = {}
    for e, nt, dt, op, ext in prepared:
        if len(op) >= 80:
            groups.setdefault(op, []).append(e["seq"])
    boiler = set()
    for op, seqs in groups.items():
        if len(seqs) >= 3:
            boiler.update(seqs)
            notes.append(f"{len(seqs)} documents share the same opening text "
                         f"({', '.join(f'[{q:04d}]' for q in sorted(seqs))}) — "
                         f"boilerplate header or a commentary series; not paired")
    pairs = []
    for i in range(len(prepared)):
        ea, ta, da, oa, xa = prepared[i]
        for j in range(i + 1, len(prepared)):
            eb, tb, db, ob, xb = prepared[j]
            reason = None
            if ta and tb and len(ta) >= 6 and da == db:
                if ta == tb:
                    reason = ("same title" if xa == xb else
                              f"same title, {xa[1:]} / {xb[1:]} pair")
                else:
                    r = difflib.SequenceMatcher(None, ta, tb).ratio()
                    if r >= title_cut:
                        reason = f"title similarity {r:.2f}"
            if (reason is None and oa and ob and len(oa) >= 80 and len(ob) >= 80
                    and ea["seq"] not in boiler and eb["seq"] not in boiler):
                r = difflib.SequenceMatcher(None, oa, ob).ratio()
                if r >= opening_cut:
                    reason = f"opening text similarity {r:.2f}"
            if reason:
                pairs.append((ea["seq"], eb["seq"], reason))
    return pairs, notes


# ------------------------------------------------------------------ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("matter_path")
    ap.add_argument("output_path")
    ap.add_argument("--recursive", action="store_true", default=True)
    ap.add_argument("--no-recursive", dest="recursive", action="store_false")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--mirror-index", default=None,
                    help="path to mirror-index.json (default .ai/mirror-index.json)")
    ap.add_argument("--exclude-dir", action="append", default=[])
    ap.add_argument("--exclude-file", action="append", default=[])
    ap.add_argument("--include-hopper-input", action="store_true",
                    help="TEST ONLY: also index 01 Hopper/Input (unfiled documents); refused under a 'live cases' path")
    ap.add_argument("--text-poor-lines", type=int, default=TEXT_POOR_LINES_PER_PAGE,
                    help="second text-poor signal: minimum text lines per page (default 2)")
    ap.add_argument("--gate", dest="gate", action="store_true", default=None,
                    help="apply the phase 2 gate (default on mirror-first matters)")
    ap.add_argument("--no-gate", dest="gate", action="store_false",
                    help="do not apply the phase 2 gate")
    ap.add_argument("--profile", default=None,
                    help="jurisdiction profile JSON (default: _intake-profile.json in the matter folder or a parent)")
    ap.add_argument("--no-versions", action="store_true",
                    help="skip the probable-versions (title / opening similarity) pass")
    args = ap.parse_args()
    PROFILE = find_profile(args.matter_path, args.profile)
    apply_profile(PROFILE)

    if args.include_hopper_input:
        parts_lower = [p.lower() for p in Path(args.matter_path).resolve().parts]
        if "live cases" in parts_lower:
            print("REFUSED: --include-hopper-input is a workflow-test flag and may not be "
                  "used on a live matter (path is under 'live cases'). Nothing written.")
            return 3

    src = Path(args.matter_path).resolve()
    out = Path(args.output_path)
    if not out.is_absolute():
        out = (Path.cwd() / out).resolve()
    docs_dir = out / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)

    excl_dirs = EXCLUDE_DIRS | set(args.exclude_dir)
    excl_files = EXCLUDE_FILES | set(args.exclude_file)

    state_file = out / "_state.json"
    state = json.loads(state_file.read_text()) if state_file.exists() else {}

    mirrors, mirror_index_path = load_mirror_index(src, args.mirror_index)

    def wanted(f):
        if not f.is_file() or f.suffix.lower() not in SUPPORTED:
            return False
        if out == f or out in f.parents:
            return False
        parts = f.relative_to(src).parts
        if any(p.startswith("_") or p.startswith(".") for p in parts):
            return False
        if args.include_hopper_input and len(parts) > 2 \
                and parts[0] == "01 Hopper" and parts[1] == "Input":
            return True
        if any(p in excl_dirs for p in parts[:-1]):
            return False
        if f.name in excl_files:
            return False
        return True

    pattern = "**/*" if args.recursive else "*"
    # "00 Current" holds copies of filed documents: walk it last so the filed
    # copy is the primary entry and the Current copy is the duplicate.
    files = sorted((f for f in src.glob(pattern) if wanted(f)),
                   key=lambda p: (p.relative_to(src).parts[0] == "00 Current",
                                  str(p).lower()))

    indexed = skipped = dupes = 0
    seq = max((v.get("seq", 0) for v in state.values()), default=0)
    seen_hash = {v["sha256"]: rel for rel, v in state.items()
                 if v.get("entry") and "sha256" in v}

    for f in files:
        rel = str(f.relative_to(src))
        digest = sha256(f)
        if not args.force and state.get(rel, {}).get("sha256") == digest:
            skipped += 1
            seen_hash.setdefault(digest, rel)
            continue
        if digest in seen_hash and seen_hash[digest] != rel:
            state[rel] = {"sha256": digest, "seq": state.get(rel, {}).get("seq", 0),
                          "status": "duplicate",
                          "duplicate_of": seen_hash[digest]}
            dupes += 1
            continue
        seq_n = state.get(rel, {}).get("seq")
        if seq_n is None:
            seq += 1
            seq_n = seq
        ext = f.suffix.lower()
        problem = None
        row = mirrors.get(digest)
        mirror_note = None
        if row:
            text, n_pages, status, problem = read_mirror(src, row, args.text_poor_lines)
            doc_ref = row["mirror_path"]
            mirror_note = (
                f"verified={row.get('verified')}, "
                f"complete_reading={row.get('complete_reading')}, "
                f"method={row.get('extraction_method', '?')}")
            if row.get("limitations"):
                mirror_note += f"; limitations: {row['limitations']}"
            if status == "failed":
                state[rel] = {"sha256": digest, "seq": seq_n, "status": "failed",
                              "problem": problem}
                continue
        else:
            try:
                text, n_pages, status = EXTRACTORS[ext](f)
            except Exception as e:
                state[rel] = {"sha256": digest, "seq": seq_n, "status": "failed",
                              "problem": f"extraction failed: {e}"}
                continue
            if status in ("scanned", "empty"):
                problem = f"{status} — needs OCR or manual review"
            docname = f"{seq_n:04d}_{safe_name(f.stem)}.md"
            (docs_dir / docname).write_text(
                f"---\nsource: {rel}\nextracted: {datetime.now():%Y-%m-%d %H:%M}\n"
                f"status: {status}\nnote: no intake mirror — extracted by "
                f"matter-index, not intake-verified\n---\n\n{text}\n")
            doc_ref = f"docs/{docname}"
            if mirror_index_path is not None:
                problem = problem or "no intake mirror — extracted by index; not intake-verified"

        typ, typ_conf = detect_type(text, ext)
        date, date_conf = detect_date(text, f.name)
        parties, parties_conf = detect_parties(text)
        deponent, dep_conf = detect_deponent(text)
        caseref, _ = detect_caseref(text)
        tribunal, _ = detect_tribunal(text)
        entry = {
            "seq": seq_n, "source": rel, "sha256": digest,
            "type": typ, "type_conf": typ_conf,
            "date": date, "date_conf": date_conf,
            "deponent": deponent, "dep_conf": dep_conf,
            "parties": parties, "parties_conf": parties_conf,
            "caseref": caseref, "tribunal": tribunal,
            "pages": n_pages, "paras": count_numbered_paras(text),
            "para_markers": len(PARA_MARKER.findall(text)),   # v2.5.1 DOCX mirrors
            "legibility": status, "doc": doc_ref,
            "mirror": mirror_note,
            "first_para": first_substantive_para(text),
        }
        state[rel] = {"sha256": digest, "seq": seq_n, "status": status,
                      "entry": entry}
        if problem:
            state[rel]["problem"] = problem
        seen_hash[digest] = rel
        indexed += 1

    existing = {str(f.relative_to(src)) for f in files}
    for gone in [k for k in state if k not in existing]:
        del state[gone]

    entries = sorted((v["entry"] for v in state.values() if "entry" in v),
                     key=lambda e: e["seq"])
    dup_rows = sorted((rel, v["duplicate_of"]) for rel, v in state.items()
                      if v.get("status") == "duplicate")
    lines = [f"# Matter index — {src.name}",
             f"*Built {datetime.now():%Y-%m-%d %H:%M}. "
             f"{len(entries)} documents, {len(dup_rows)} duplicates. "
             f"Text source: {'intake mirrors (' + str(mirror_index_path.relative_to(src)) + ')' if mirror_index_path else 'index extraction'}. "
             f"Heuristic metadata — verify low/medium confidence fields against the source. "
             f"Cite the source file and physical page, never the mirror filename.*", ""]
    for e in entries:
        title = Path(e["source"]).stem
        lines += [
            f"## [{e['seq']:04d}] {title}",
            "",
            f"- **Source:** `{e['source']}`",
            f"- **SHA-256:** `{e['sha256'][:16]}…`",
            f"- **Type:** {e['type']} *(confidence: {e['type_conf']})*",
            f"- **Date:** {e['date']} *(confidence: {e['date_conf']})*",
            f"- **Author / Deponent:** {e['deponent']} *(confidence: {e['dep_conf']})*",
            f"- **Parties:** {e['parties']} *(confidence: {e['parties_conf']})*",
            f"- **Case / matter ref:** {e['caseref']}",
            f"- **Tribunal / court:** {e['tribunal']}",
            f"- **Page count:** {e['pages'] if e['pages'] else '—'}"
            + (f" *(docx mirror: {e['para_markers']} paragraph markers — cite as \"(docx, para n)\")*"
               if e.get('para_markers') else ""),
            f"- **Numbered paragraphs detected:** ~ {e['paras']}",
            f"- **Legibility / text status:** {e['legibility']}",
            f"- **Reading copy:** `{e['doc']}`",
        ]
        if e.get("mirror"):
            lines.append(f"- **Intake mirror status:** {e['mirror']}")
        lines.append("")
        if e["first_para"]:
            lines += ["**First substantive paragraph:**", "",
                      f"> {e['first_para']}", ""]
        lines.append("---")
        lines.append("")
    if dup_rows:
        lines += ["## Duplicates (same SHA-256, not indexed twice)", ""]
        lines += [f"- `{rel}` = `{of}`" for rel, of in dup_rows]
        lines.append("")
    pv_pairs, pv_notes = ([], ["skipped (--no-versions)"]) if args.no_versions \
        else probable_versions(entries)
    by_seq = {e["seq"]: e for e in entries}
    lines += ["## Probable versions (title or opening-text similarity — verify by reading)", ""]
    for n in pv_notes:
        lines.append(f"- {n}")
    if pv_notes:
        lines.append("")
    if pv_pairs:
        lines.append("*Heuristic. Each pair may be the same document in another file or "
                     "format (docx/pdf, two scans, draft v signed) or two different "
                     "documents with alike names. The map readers confirm; nothing here "
                     "is a fact about the documents.*")
        lines.append("")
        for a, b, why in pv_pairs:
            lines.append(f"- [{a:04d}] `{by_seq[a]['source']}` ~ [{b:04d}] "
                         f"`{by_seq[b]['source']}` — {why}")
    elif not args.no_versions:
        lines.append("None found.")
    lines.append("")
    (out / "index.md").write_text("\n".join(lines))

    ill_lines = ["# Illegible / problem files", ""]
    problems = sorted(
        (rel, v["problem"]) for rel, v in state.items() if v.get("problem"))
    if problems:
        for rel, why in problems:
            ill_lines.append(f"- `{rel}` — {why}")
    else:
        ill_lines.append("None.")
    (out / "_illegible.md").write_text("\n".join(ill_lines) + "\n")

    state_file.write_text(json.dumps(state, indent=1))
    print(f"PASS: indexed {indexed}, skipped (unchanged) {skipped}, "
          f"duplicates {dupes}, problems {len(problems)}, "
          f"probable versions {len(pv_pairs)}, "
          f"mirror index: {mirror_index_path or 'none'}")
    print(f"index: {out / 'index.md'}")

    # v2.5 phase 2 gate: filed documents > 0 and text-poor mirrors = 0,
    # else stop and report. Default on for mirror-first matters.
    gate_on = args.gate if args.gate is not None else (mirror_index_path is not None)
    if gate_on:
        text_poor = sorted(rel for rel, v in state.items()
                           if v.get("status") == "mirror-text-poor")
        n_docs = len(entries)
        if n_docs == 0 or text_poor:
            print(f"GATE: FAIL — filed documents indexed {n_docs}, "
                  f"text-poor mirrors {len(text_poor)}. Stop and report; do not "
                  f"start the matter map on this index without the user's say-so.")
            for rel in text_poor[:50]:
                print(f"  text-poor: {rel}")
            return 2
        print(f"GATE: PASS — {n_docs} filed documents, 0 text-poor mirrors.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
